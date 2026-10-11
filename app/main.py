
import os
import base64
import re
import uuid
import logging
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from .db import init_db, add_message, get_messages, add_memory, get_memories
from .prompts import SYSTEM_PROMPT

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# بارگذاری تنظیمات از پوشه اصلی پروژه
BASE_DIR = Path(__file__).resolve().parent.parent
load_dotenv(BASE_DIR / ".env", override=True)

API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
MODEL = (os.getenv("OPENAI_MODEL") or "gpt-4.1").strip()
MAX_HISTORY = int(os.getenv("MAX_HISTORY", "30"))
APP_NAME = os.getenv("APP_NAME", "یار هوشمند")

app = FastAPI(title=APP_NAME, version="1.0.1")
app.mount(
    "/static",
    StaticFiles(directory=BASE_DIR / "static"),
    name="static",
)

class ChatRequest(BaseModel):
    session_id: str
    message: str
    mode: str = "auto"
    web_search: bool = True
    remember: bool = False

class MemoryRequest(BaseModel):
    session_id: str
    memory: str

MODE_HINTS = {
    "auto": "خودت بهترین روش پاسخ را انتخاب کن.",
    "general": "حالت عمومی.",
    "medical": "حالت سلامت و پزشکی با رعایت محدودیت‌ها.",
    "research": "حالت تحقیق؛ منابع و عدم قطعیت را جدی بگیر.",
    "coding": "حالت برنامه‌نویسی؛ کد اجرایی و توضیح دقیق بده.",
    "finance": "حالت مالی؛ ریسک و عدم قطعیت را شفاف کن.",
    "writing": "حالت نویسندگی؛ متن نهایی و قابل استفاده تولید کن.",
}

def get_client():
    if not API_KEY:
        raise HTTPException(
            status_code=503,
            detail="کلید API تنظیم نشده است. فایل .env را بررسی کن.",
        )

    # جلوگیری از ارسال کلید نمونه یا کلید دارای حروف فارسی
    placeholders = (
        "your_api_key_here",
        "کلید_واقعی_خودت",
        "کلیدواقعی",
        "اینجا_قرار_بده",
    )
    if not API_KEY.isascii() or any(p in API_KEY for p in placeholders):
        raise HTTPException(
            status_code=503,
            detail=(
                "کلید API معتبر وارد نشده است. در فایل .env "
                "کلید واقعی انگلیسی را قرار بده؛ متن نمونه قابل استفاده نیست."
            ),
        )

    return OpenAI(api_key=API_KEY)

def extract_file_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(data)
            tmp = Path(f.name)
        try:
            reader = PdfReader(str(tmp))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            )[:50000]
        finally:
            tmp.unlink(missing_ok=True)

    if ext == ".docx":
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(data)
            tmp = Path(f.name)
        try:
            doc = Document(str(tmp))
            return "\n".join(p.text for p in doc.paragraphs)[:50000]
        finally:
            tmp.unlink(missing_ok=True)

    if ext == ".xlsx":
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(data)
            tmp = Path(f.name)
        try:
            wb = load_workbook(tmp, read_only=True, data_only=True)
            chunks = []
            try:
                for ws in wb.worksheets:
                    chunks.append(f"[Sheet: {ws.title}]")
                    for row in ws.iter_rows(values_only=True):
                        chunks.append(
                            " | ".join(
                                "" if value is None else str(value)
                                for value in row
                            )
                        )
                        if sum(map(len, chunks)) > 50000:
                            break
            finally:
                wb.close()
            return "\n".join(chunks)[:50000]
        finally:
            tmp.unlink(missing_ok=True)

    if ext in {".txt", ".md", ".csv", ".json"}:
        return data.decode("utf-8-sig", errors="replace")[:50000]

    return ""

def build_input(session_id: str, message: str, mode: str, memories):
    history = get_messages(session_id, MAX_HISTORY)
    history_text = [
        f"{item['role']}: {item['content']}" for item in history
    ]
    memory_text = "\n".join(f"- {m}" for m in memories) or "حافظه‌ای ثبت نشده."
    context = (
        f"حالت انتخاب‌شده: {MODE_HINTS.get(mode, MODE_HINTS['auto'])}\n"
        f"حافظه مرتبط:\n{memory_text}\n\n"
        f"سابقه اخیر:\n" + "\n".join(history_text[-MAX_HISTORY:])
    )
    return [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "system", "content": context},
        {"role": "user", "content": message},
    ]

@app.on_event("startup")
def startup():
    init_db()

@app.get("/")
def index():
    return FileResponse(BASE_DIR / "static" / "index.html")

@app.get("/api/health")
def health():
    return {
        "ok": True,
        "configured": bool(API_KEY and API_KEY.isascii()),
        "model": MODEL,
    }

@app.get("/api/memories/{session_id}")
def list_memories(session_id: str):
    return {"memories": get_memories(session_id)}

@app.post("/api/memories")
def create_memory(req: MemoryRequest):
    memory = req.memory.strip()
    if not memory:
        raise HTTPException(400, "حافظه خالی است.")
    add_memory(req.session_id, memory[:1000])
    return {"ok": True}

@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "حداکثر اندازه فایل ۱۰ مگابایت است.")

    filename = file.filename or "file"
    text = extract_file_text(filename, data)
    if text:
        return {"kind": "text", "filename": filename, "text": text}

    if (file.content_type or "").startswith("image/"):
        return {
            "kind": "image",
            "filename": filename,
            "mime": file.content_type,
            "data": base64.b64encode(data).decode("ascii"),
        }

    raise HTTPException(415, "این نوع فایل پشتیبانی نمی‌شود.")

@app.post("/api/chat")
def chat(req: ChatRequest):
    message = req.message.strip()
    if not message:
        raise HTTPException(400, "پیام خالی است.")

    # بررسی کلید پیش از تماس با سرویس
    client = get_client()

    memories = get_memories(req.session_id)

    if req.remember or re.search(
        r"(یادت باشه|یاد بگیر|به خاطر بسپار)", message
    ):
        cleaned = re.sub(
            r"^.*?(یادت باشه|یاد بگیر|به خاطر بسپار)\s*",
            "",
            message,
        ).strip(" :،")
        if cleaned:
            add_memory(req.session_id, cleaned[:1000])
            memories = get_memories(req.session_id)

    kwargs = {
        "model": MODEL,
        "input": build_input(req.session_id, message, req.mode, memories),
    }

    if req.web_search:
        kwargs["tools"] = [{"type": "web_search"}]

    try:
        response = client.responses.create(**kwargs)
        answer = response.output_text
    except Exception as exc:
        # جزئیات فقط در ترمینال؛ پاسخ امن‌تر برای کاربر
        logger.exception("OpenAI request failed")
        error_name = type(exc).__name__
        raise HTTPException(
            status_code=502,
            detail=f"ارتباط با مدل ناموفق بود ({error_name}). "
                   "خطای کامل در ترمینال ثبت شده است.",
        )

    add_message(req.session_id, "user", message)
    add_message(req.session_id, "assistant", answer)

    return {
        "answer": answer,
        "session_id": req.session_id,
        "model": MODEL,
    }

@app.get("/api/history/{session_id}")
def history(session_id: str):
    return {"messages": get_messages(session_id, MAX_HISTORY)}
